"""R04 mechanism and decision diagnostics for the SAVP manuscript.

This script deliberately works from the existing prediction-level exports and
inspectable public feature tables. It adds auditable diagnostics that sharpen
the R02 evidence without changing the underlying benchmark scores:

1. topology-preserving pseudo-group resampling nulls;
2. group-shift fingerprints linking held-out group stress to errors;
3. worst-group / CVaR-style decision-risk summaries;
4. screening-regret summaries for capacity-like regression rows;
5. residual/probability calibration diagnostics;
6. source-balanced scoring probes.

The source-balanced probe is an evaluation-weighting diagnostic, not a retrained
source-balanced ERM benchmark. A retrained mitigation baseline would require a
separate per-module QA cycle.
"""

from __future__ import annotations

import csv
import math
import random
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, median
from typing import Iterable

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional dependency in reproduction envs
    load_workbook = None  # type: ignore[assignment]


ROOT = Path(__file__).resolve().parents[1]
REPRO = ROOT / "code" / "outputs" / "reproductions"
DATASETS = ROOT / "code" / "outputs" / "datasets"
FIG_OUT = ROOT / "code" / "outputs" / "figures"
TABLE_OUT = ROOT / "manuscript" / "tables"
RANDOM_SEED = 20260513
NULL_REPS = 1000
EPS = 1e-12


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def f(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "" or text.lower() in {"nan", "none", "null"}:
        return None
    try:
        out = float(text)
    except ValueError:
        return None
    if not math.isfinite(out):
        return None
    return out


def fmt(value: float | None, digits: int = 3) -> str:
    if value is None or not math.isfinite(value):
        return ""
    return f"{value:.{digits}f}"


def clean_text(text: str) -> str:
    replacements = {
        "–": "-",
        "—": "-",
        "−": "-",
        "×": "x",
        "²": "2",
        "Δ": "Delta",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return " ".join(text.split())


def model_name(value: str) -> str:
    text = value.replace("_", " ").replace("-", " ").strip()
    compact = text.lower().replace(" ", "")
    if compact in {"randomforest", "randomforestregressor", "randomforestclassifier"}:
        return "Random Forest"
    if compact in {"gradientboosting", "gradientboostingregressor", "gradientboostingclassifier"}:
        return "Gradient Boosting"
    if compact in {"logisticregression", "logistic"}:
        return "Logistic Regression"
    if compact in {"ridge", "ridgeregression"}:
        return "Ridge"
    return text or value


def split_kind(row: dict[str, str]) -> str | None:
    value = row.get("validation", row.get("split", row.get("cv", "")))
    if "Random" in value:
        return "random"
    if "GroupKFold" in value:
        return "grouped"
    return None


def prediction_descriptor(module: str, row: dict[str, str]) -> str:
    parts = [row.get(key, "") for key in ("case", "task", "target") if row.get(key, "")]
    return clean_text("|".join(parts)) if parts else module


def prediction_columns(row: dict[str, str]) -> tuple[str, str, str, str]:
    y_col = "observed" if "observed" in row else "y_true"
    p_col = "predicted" if "predicted" in row else "y_pred"
    id_col = "sample_index" if "sample_index" in row else "row_index"
    if "source_group" in row:
        group_col = "source_group"
    elif "mix_group" in row:
        group_col = "mix_group"
    elif "shape" in row:
        group_col = "shape"
    else:
        group_col = "group"
    return y_col, p_col, id_col, group_col


def normalize_label(value: object) -> str:
    text = str(value).strip()
    number = f(text)
    if number is not None and abs(number - round(number)) < 1e-9:
        return str(int(round(number)))
    return text


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    xs = sorted(values)
    if len(xs) == 1:
        return xs[0]
    pos = q * (len(xs) - 1)
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return xs[lo]
    return xs[lo] * (hi - pos) + xs[hi] * (pos - lo)


def ranks(values: list[float]) -> list[float]:
    indexed = sorted(enumerate(values), key=lambda item: item[1])
    out = [0.0] * len(values)
    i = 0
    while i < len(indexed):
        j = i + 1
        while j < len(indexed) and indexed[j][1] == indexed[i][1]:
            j += 1
        rank = (i + j + 1) / 2.0
        for k in range(i, j):
            out[indexed[k][0]] = rank
        i = j
    return out


def pearson(x: list[float], y: list[float]) -> float | None:
    if len(x) != len(y) or len(x) < 3:
        return None
    mx = mean(x)
    my = mean(y)
    sx = math.sqrt(sum((v - mx) ** 2 for v in x))
    sy = math.sqrt(sum((v - my) ** 2 for v in y))
    if sx <= EPS or sy <= EPS:
        return None
    return sum((a - mx) * (b - my) for a, b in zip(x, y)) / (sx * sy)


def spearman(x: list[float], y: list[float]) -> float | None:
    if len(x) != len(y) or len(x) < 3:
        return None
    return pearson(ranks(x), ranks(y))


def variance(values: list[float]) -> float | None:
    if len(values) < 2:
        return None
    m = mean(values)
    return sum((value - m) ** 2 for value in values) / (len(values) - 1)


def r2_score(y: list[float], pred: list[float]) -> float | None:
    if len(y) < 2 or len(y) != len(pred):
        return None
    y_mean = mean(y)
    ss_tot = sum((value - y_mean) ** 2 for value in y)
    if ss_tot <= EPS:
        return None
    ss_res = sum((value - guess) ** 2 for value, guess in zip(y, pred))
    return 1.0 - ss_res / ss_tot


def mae(y: list[float], pred: list[float]) -> float | None:
    if not y or len(y) != len(pred):
        return None
    return mean(abs(value - guess) for value, guess in zip(y, pred))


def accuracy(y: list[str], pred: list[str]) -> float | None:
    if not y or len(y) != len(pred):
        return None
    return sum(1 for yy, pp in zip(y, pred) if yy == pp) / len(y)


def balanced_accuracy(y: list[str], pred: list[str]) -> float | None:
    if not y or len(y) != len(pred):
        return None
    labels = sorted(set(y))
    recalls: list[float] = []
    for label in labels:
        idx = [i for i, yy in enumerate(y) if yy == label]
        if idx:
            recalls.append(sum(1 for i in idx if pred[i] == label) / len(idx))
    return mean(recalls) if recalls else None


def macro_f1(y: list[str], pred: list[str]) -> float | None:
    if not y or len(y) != len(pred):
        return None
    labels = sorted(set(y) | set(pred))
    scores: list[float] = []
    for label in labels:
        tp = sum(1 for yy, pp in zip(y, pred) if yy == label and pp == label)
        fp = sum(1 for yy, pp in zip(y, pred) if yy != label and pp == label)
        fn = sum(1 for yy, pp in zip(y, pred) if yy == label and pp != label)
        denom = 2 * tp + fp + fn
        scores.append((2 * tp / denom) if denom else 0.0)
    return mean(scores) if scores else None


def score_classification(y: list[str], pred: list[str], metric: str = "accuracy") -> float | None:
    if metric == "accuracy":
        return accuracy(y, pred)
    if metric == "balanced_accuracy":
        return balanced_accuracy(y, pred)
    if metric == "macro_f1":
        return macro_f1(y, pred)
    raise ValueError(metric)


def scale_for(y: list[float]) -> float:
    return mean(abs(value) for value in y) or 1.0


def is_classification_descriptor(descriptor: str, rows: list[dict[str, str]]) -> bool:
    text = descriptor.lower()
    if "classification" in text or "failure-mode" in text or "failure mode" in text:
        return True
    if rows and any(key.startswith("prob_class_") for key in rows[0]):
        return True
    return False


def collect_prediction_pairs() -> list[dict[str, object]]:
    paired_sets: list[dict[str, object]] = []
    for path in sorted(REPRO.glob("*/predictions.csv")):
        module = path.parent.name
        rows = read_csv(path)
        buckets: dict[tuple[str, str], dict[str, dict[str, dict[str, str]]]] = defaultdict(lambda: defaultdict(dict))
        group_labels: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
        for row in rows:
            kind = split_kind(row)
            if kind is None:
                continue
            _, _, id_col, group_col = prediction_columns(row)
            descriptor = prediction_descriptor(module, row)
            key = (descriptor, model_name(row.get("model", "")))
            sample_id = row.get(id_col, "")
            if sample_id == "":
                continue
            buckets[key][kind][sample_id] = row
            if kind == "grouped":
                group_labels[key][sample_id] = row.get(group_col, "")
        for (descriptor, model), split_rows in buckets.items():
            if "random" not in split_rows or "grouped" not in split_rows:
                continue
            common = sorted(set(split_rows["random"]) & set(split_rows["grouped"]), key=lambda x: (f(x) is None, f(x) or 0, x))
            if len(common) < 10:
                continue
            random_rows = [split_rows["random"][sample] for sample in common]
            grouped_rows = [split_rows["grouped"][sample] for sample in common]
            y_col, p_col, _, _ = prediction_columns(grouped_rows[0])
            classification = is_classification_descriptor(descriptor, grouped_rows)
            if not classification:
                numeric = all(f(row.get(y_col)) is not None and f(row.get(p_col)) is not None for row in random_rows + grouped_rows)
                classification = not numeric
            paired_sets.append(
                {
                    "module": module,
                    "descriptor": descriptor,
                    "model": model,
                    "metric_type": "classification" if classification else "regression",
                    "sample_ids": common,
                    "random_rows": random_rows,
                    "grouped_rows": grouped_rows,
                    "group_labels": [group_labels[(descriptor, model)].get(sample, "ungrouped") for sample in common],
                    "y_col": y_col,
                    "p_col": p_col,
                }
            )
    return paired_sets


def numeric_vectors(data: dict[str, object]) -> tuple[list[float], list[float], list[float]]:
    random_rows = data["random_rows"]  # type: ignore[index]
    grouped_rows = data["grouped_rows"]  # type: ignore[index]
    y_col = str(data["y_col"])
    p_col = str(data["p_col"])
    y: list[float] = []
    random_pred: list[float] = []
    grouped_pred: list[float] = []
    for rr, gr in zip(random_rows, grouped_rows):  # type: ignore[union-attr]
        yy = f(rr.get(y_col))
        pr = f(rr.get(p_col))
        pg = f(gr.get(p_col))
        if yy is None or pr is None or pg is None:
            continue
        y.append(yy)
        random_pred.append(pr)
        grouped_pred.append(pg)
    return y, random_pred, grouped_pred


def class_vectors(data: dict[str, object]) -> tuple[list[str], list[str], list[str]]:
    random_rows = data["random_rows"]  # type: ignore[index]
    grouped_rows = data["grouped_rows"]  # type: ignore[index]
    y_col = str(data["y_col"])
    p_col = str(data["p_col"])
    y = [normalize_label(row.get(y_col)) for row in random_rows]  # type: ignore[union-attr]
    random_pred = [normalize_label(row.get(p_col)) for row in random_rows]  # type: ignore[union-attr]
    grouped_pred = [normalize_label(row.get(p_col)) for row in grouped_rows]  # type: ignore[union-attr]
    return y, random_pred, grouped_pred


def index_groups(labels: list[str]) -> dict[str, list[int]]:
    groups: dict[str, list[int]] = defaultdict(list)
    for idx, label in enumerate(labels):
        groups[str(label)].append(idx)
    return groups


def group_sizes_from(labels: list[str]) -> list[int]:
    return sorted((len(indices) for indices in index_groups(labels).values()), reverse=True)


def shuffled_group_labels(labels: list[str], rng: random.Random) -> list[str]:
    sizes = group_sizes_from(labels)
    indices = list(range(len(labels)))
    rng.shuffle(indices)
    out = [""] * len(labels)
    cursor = 0
    for gid, size in enumerate(sizes):
        for idx in indices[cursor : cursor + size]:
            out[idx] = f"pseudo_{gid:03d}"
        cursor += size
    return out


def group_balanced_mean(values: list[float], labels: list[str]) -> float | None:
    groups = index_groups(labels)
    means: list[float] = []
    for idx in groups.values():
        vals = [values[i] for i in idx]
        if vals:
            means.append(mean(vals))
    return mean(means) if means else None


def topology_pseudogroup_null(paired_sets: list[dict[str, object]]) -> None:
    rng = random.Random(RANDOM_SEED)
    rows: list[dict[str, object]] = []
    for data in paired_sets:
        labels = [str(x) for x in data["group_labels"]]  # type: ignore[index]
        n_groups = len(set(labels))
        if n_groups < 2:
            continue
        if data["metric_type"] == "regression":
            y, pr, pg = numeric_vectors(data)
            if len(y) != len(labels):
                continue
            scale = scale_for(y)
            paired_loss_delta = [(abs(yy - gg) - abs(yy - rr)) / scale for yy, rr, gg in zip(y, pr, pg)]
            pooled_random = r2_score(y, pr)
            pooled_grouped = r2_score(y, pg)
            pooled_gap = None if pooled_random is None or pooled_grouped is None else pooled_random - pooled_grouped
            basis = "group-balanced normalized absolute-error inflation"
        else:
            y_cls, pr_cls, pg_cls = class_vectors(data)
            if len(y_cls) != len(labels):
                continue
            paired_loss_delta = [
                (0.0 if yy == gg else 1.0) - (0.0 if yy == rr else 1.0)
                for yy, rr, gg in zip(y_cls, pr_cls, pg_cls)
            ]
            pooled_random = accuracy(y_cls, pr_cls)
            pooled_grouped = accuracy(y_cls, pg_cls)
            pooled_gap = None if pooled_random is None or pooled_grouped is None else pooled_random - pooled_grouped
            basis = "group-balanced classification error inflation"
        actual = group_balanced_mean(paired_loss_delta, labels)
        if actual is None:
            continue
        null_values: list[float] = []
        for _ in range(NULL_REPS):
            pseudo_labels = shuffled_group_labels(labels, rng)
            value = group_balanced_mean(paired_loss_delta, pseudo_labels)
            if value is not None and math.isfinite(value):
                null_values.append(value)
        lo = percentile(null_values, 0.025)
        hi = percentile(null_values, 0.975)
        med = percentile(null_values, 0.50)
        percentile_rank = sum(1 for value in null_values if value <= actual) / len(null_values) if null_values else None
        if hi is not None and actual > hi:
            interpretation = "error_concentrates_beyond_topology_null"
        elif lo is not None and actual < lo:
            interpretation = "actual_below_topology_null"
        else:
            interpretation = "compatible_with_topology_null"
        rows.append(
            {
                "module": data["module"],
                "descriptor": data["descriptor"],
                "model": data["model"],
                "metric_type": data["metric_type"],
                "n_samples": len(labels),
                "n_groups": n_groups,
                "median_group_size": fmt(median(group_sizes_from(labels)), 2),
                "pooled_random_score": fmt(pooled_random, 6),
                "pooled_grouped_score": fmt(pooled_grouped, 6),
                "pooled_random_minus_grouped_gap": fmt(pooled_gap, 6),
                "null_metric_basis": basis,
                "actual_group_balanced_loss_gap": fmt(actual, 6),
                "pseudo_null_median": fmt(med, 6),
                "pseudo_null_95_low": fmt(lo, 6),
                "pseudo_null_95_high": fmt(hi, 6),
                "actual_percentile_in_null": fmt(percentile_rank, 3),
                "actual_minus_null_median": fmt(None if med is None else actual - med, 6),
                "null_reps": len(null_values),
                "interpretation": interpretation,
            }
        )
    fields = [
        "module",
        "descriptor",
        "model",
        "metric_type",
        "n_samples",
        "n_groups",
        "median_group_size",
        "pooled_random_score",
        "pooled_grouped_score",
        "pooled_random_minus_grouped_gap",
        "null_metric_basis",
        "actual_group_balanced_loss_gap",
        "pseudo_null_median",
        "pseudo_null_95_low",
        "pseudo_null_95_high",
        "actual_percentile_in_null",
        "actual_minus_null_median",
        "null_reps",
        "interpretation",
    ]
    write_csv(FIG_OUT / "r04_topology_pseudogroup_null.csv", rows, fields)

    counts = Counter(str(row["interpretation"]) for row in rows)
    lines = [
        "# Table S15. R04 topology-preserving pseudo-group null diagnostic",
        "",
        "This is a prediction-level pseudo-group resampling diagnostic. It preserves the grouped-validation group-size distribution but randomly reassigns samples to pseudo-groups, then recomputes group-balanced paired loss inflation. It does not retrain GroupKFold models under pseudo-group assignments.",
        "",
        f"- Descriptor/model rows analysed: {len(rows)}.",
        f"- Rows where true grouped errors concentrate beyond the topology null: {counts.get('error_concentrates_beyond_topology_null', 0)}.",
        "",
        "| Descriptor | model | basis | actual | pseudo-null 95% interval | percentile | interpretation |",
        "|---|---|---|---:|---|---:|---|",
    ]
    for row in rows:
        if row["model"] == "Random Forest" or row["metric_type"] == "classification":
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['null_metric_basis']} | {row['actual_group_balanced_loss_gap']} | {row['pseudo_null_95_low']} to {row['pseudo_null_95_high']} | {row['actual_percentile_in_null']} | {row['interpretation']} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: this table separates error concentration over the declared deployment groups from group-size geometry alone. It is a mechanism diagnostic, not a prevalence estimate or a causal topology model.",
        ]
    )
    write_text(TABLE_OUT / "r04_topology_pseudogroup_null.md", "\n".join(lines) + "\n")


def csv_feature_map(path: Path, sample_id_col: str | None, exclude: set[str]) -> dict[str, dict[str, float]]:
    if not path.exists():
        return {}
    rows = read_csv(path)
    out: dict[str, dict[str, float]] = {}
    for ordinal, row in enumerate(rows):
        sample_id = str(row.get(sample_id_col, ordinal)) if sample_id_col else str(ordinal)
        features: dict[str, float] = {}
        for key, value in row.items():
            if key in exclude:
                continue
            number = f(value)
            if number is not None:
                features[key] = number
        if features:
            out[sample_id] = features
    return out


def mangalathu_feature_map() -> dict[str, dict[str, float]]:
    if load_workbook is None:
        return {}
    path = DATASETS / "mangalathu_2020_shear_wall" / "Shear_Wall_Database.xlsx"
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(value).strip() if value is not None else f"unnamed_{idx}" for idx, value in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))]
    exclude = {"unnamed_0", "Author", "Specimen", "FailureMode", "Section"}
    out: dict[str, dict[str, float]] = {}
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        features: dict[str, float] = {}
        for key, value in zip(header, row):
            if key in exclude:
                continue
            number = f(value)
            if number is not None:
                features[key] = number
        if features:
            out[str(idx)] = features
    return out


def sfrc_feature_map() -> dict[str, dict[str, float]]:
    if load_workbook is None:
        return {}
    path = DATASETS / "lantsoght_2019_sfrc_shear" / "database_SFRC.xlsx"
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 4:
        return {}
    data_rows = rows[3:]
    # Positions mirror 08_reproduce_rahman_2021_sfrc_regression.py.
    positions = {
        "rho_l": 16,
        "a_over_d": 18,
        "fc": 21,
        "Vf": 25,
        "L_over_d": 27,
        "F": 29,
    }
    out: dict[str, dict[str, float]] = {}
    kept = 0
    for row in data_rows:
        features: dict[str, float] = {}
        for name, pos in positions.items():
            if pos < len(row):
                number = f(row[pos])
                if number is not None:
                    features[name] = number
        if len(features) == len(positions):
            out[str(kept)] = features
            kept += 1
    return out


def feature_map_for_module(module: str) -> dict[str, dict[str, float]]:
    if module == "10-1016-j-engstruct-2019-110331":
        return mangalathu_feature_map()
    if module == "10-1016-j-engstruct-2020-111743":
        return sfrc_feature_map()
    if module == "10-1038-s41598-024-53352-1":
        return csv_feature_map(
            REPRO / module / "combined_stub_cfst_data.csv",
            None,
            {"shape", "P", "nominal_capacity", "strength_index"},
        )
    if module == "10-5281-zenodo-8062007":
        return csv_feature_map(
            REPRO / module / "analysis_data.csv",
            None,
            {"source_group", "log_mmax_exp", "Cross-section", "Test Type and Configuration"},
        )
    if module == "designsafe_prj2430_wall":
        return csv_feature_map(
            REPRO / module / "extracted_wall_data.csv",
            "row_index",
            {"row_index", "Authors", "SpecimenID", "UniqueID", "Vmax", "Dmax", "driftCap", "dispCap"},
        )
    if module == "designsafe_prj3053_coupling_beams":
        return csv_feature_map(
            REPRO / module / "extracted_coupling_beam_data.csv",
            "row_index",
            {
                "row_index",
                "reference",
                "specimen_id",
                "source_group",
                "V_m_minus_kips",
                "V_m_plus_kips",
                "V_m_avg_kips",
                "v_m_normalized",
                "CR_capacity_minus",
                "CR_capacity_plus",
                "CR_capacity_average",
            },
        )
    return {}


def feature_shift_for_group(feature_rows: dict[str, dict[str, float]], sample_ids: list[str], group_indices: list[int]) -> tuple[float | None, float | None, int]:
    group_ids = {sample_ids[i] for i in group_indices}
    available = [sid for sid in sample_ids if sid in feature_rows]
    if len(available) < 10:
        return None, None, 0
    names = sorted(set().union(*(feature_rows[sid].keys() for sid in available)))
    mean_smds: list[float] = []
    max_smds: list[float] = []
    for name in names:
        all_values = [feature_rows[sid][name] for sid in available if name in feature_rows[sid]]
        group_values = [feature_rows[sid][name] for sid in available if sid in group_ids and name in feature_rows[sid]]
        rest_values = [feature_rows[sid][name] for sid in available if sid not in group_ids and name in feature_rows[sid]]
        if len(group_values) < 1 or len(rest_values) < 2 or len(all_values) < 3:
            continue
        sd = math.sqrt(variance(all_values) or 0.0)
        if sd <= EPS:
            continue
        mean_smds.append(abs(mean(group_values) - mean(rest_values)) / sd)
    if not mean_smds:
        return None, None, 0
    max_smds.append(max(mean_smds))
    return mean(mean_smds), max(max_smds), len(mean_smds)


def entropy_from_counts(counts: Iterable[int]) -> float:
    counts = [c for c in counts if c > 0]
    total = sum(counts)
    if total == 0:
        return 0.0
    return -sum((c / total) * math.log(c / total, 2) for c in counts)


def group_shift_fingerprint(paired_sets: list[dict[str, object]]) -> None:
    group_rows: list[dict[str, object]] = []
    summary_rows: list[dict[str, object]] = []
    feature_cache: dict[str, dict[str, dict[str, float]]] = {}
    for data in paired_sets:
        labels = [str(x) for x in data["group_labels"]]  # type: ignore[index]
        groups = index_groups(labels)
        sample_ids = [str(x) for x in data["sample_ids"]]  # type: ignore[index]
        if data["module"] not in feature_cache:
            feature_cache[str(data["module"])] = feature_map_for_module(str(data["module"]))
        feature_rows = feature_cache[str(data["module"])]
        if data["metric_type"] == "regression":
            y, pr, pg = numeric_vectors(data)
            if len(y) != len(labels):
                continue
            global_sd = math.sqrt(variance(y) or 0.0) or 1.0
            scale = scale_for(y)
            for group, idx in groups.items():
                rest = [i for i in range(len(y)) if i not in set(idx)]
                if not rest:
                    continue
                grouped_loss = mean(abs(y[i] - pg[i]) / scale for i in idx)
                random_loss = mean(abs(y[i] - pr[i]) / scale for i in idx)
                target_shift = abs(mean(y[i] for i in idx) - mean(y[i] for i in rest)) / global_sd
                group_var = variance([y[i] for i in idx])
                rest_var = variance([y[i] for i in rest])
                var_ratio = None if group_var is None or rest_var is None or rest_var <= EPS else group_var / rest_var
                feat_mean, feat_max, n_feat = feature_shift_for_group(feature_rows, sample_ids, idx)
                group_rows.append(
                    {
                        "module": data["module"],
                        "descriptor": data["descriptor"],
                        "model": data["model"],
                        "metric_type": data["metric_type"],
                        "group": group,
                        "group_n": len(idx),
                        "feature_mean_smd": fmt(feat_mean, 6),
                        "feature_max_smd": fmt(feat_max, 6),
                        "n_features_used": n_feat,
                        "target_shift_smd": fmt(target_shift, 6),
                        "target_variance_ratio": fmt(var_ratio, 6),
                        "class_majority_share": "",
                        "class_entropy": "",
                        "random_group_loss": fmt(random_loss, 6),
                        "grouped_group_loss": fmt(grouped_loss, 6),
                        "paired_loss_gap": fmt(grouped_loss - random_loss, 6),
                    }
                )
        else:
            y_cls, pr_cls, pg_cls = class_vectors(data)
            if len(y_cls) != len(labels):
                continue
            labels_all = sorted(set(y_cls))
            for group, idx in groups.items():
                counts = Counter(y_cls[i] for i in idx)
                rest = [i for i in range(len(y_cls)) if i not in set(idx)]
                rest_counts = Counter(y_cls[i] for i in rest)
                total = len(idx)
                majority_share = max(counts.values()) / total if total else None
                entropy = entropy_from_counts(counts.values())
                class_l1 = 0.0
                for label in labels_all:
                    group_p = counts[label] / total if total else 0.0
                    rest_p = rest_counts[label] / len(rest) if rest else 0.0
                    class_l1 += abs(group_p - rest_p)
                random_loss = mean(0.0 if y_cls[i] == pr_cls[i] else 1.0 for i in idx)
                grouped_loss = mean(0.0 if y_cls[i] == pg_cls[i] else 1.0 for i in idx)
                feat_mean, feat_max, n_feat = feature_shift_for_group(feature_rows, sample_ids, idx)
                group_rows.append(
                    {
                        "module": data["module"],
                        "descriptor": data["descriptor"],
                        "model": data["model"],
                        "metric_type": data["metric_type"],
                        "group": group,
                        "group_n": len(idx),
                        "feature_mean_smd": fmt(feat_mean, 6),
                        "feature_max_smd": fmt(feat_max, 6),
                        "n_features_used": n_feat,
                        "target_shift_smd": fmt(class_l1 / 2.0, 6),
                        "target_variance_ratio": "",
                        "class_majority_share": fmt(majority_share, 6),
                        "class_entropy": fmt(entropy, 6),
                        "random_group_loss": fmt(random_loss, 6),
                        "grouped_group_loss": fmt(grouped_loss, 6),
                        "paired_loss_gap": fmt(grouped_loss - random_loss, 6),
                    }
                )

    by_key: dict[tuple[str, str, str, str], list[dict[str, object]]] = defaultdict(list)
    for row in group_rows:
        key = (str(row["module"]), str(row["descriptor"]), str(row["model"]), str(row["metric_type"]))
        by_key[key].append(row)
    for (module, descriptor, model, metric_type), rows in sorted(by_key.items()):
        losses = [f(row["grouped_group_loss"]) for row in rows]
        gaps = [f(row["paired_loss_gap"]) for row in rows]
        target_shift = [f(row["target_shift_smd"]) for row in rows]
        feature_shift = [f(row["feature_mean_smd"]) for row in rows]
        group_n = [f(row["group_n"]) for row in rows]
        usable_loss = [float(v) for v in losses if v is not None]
        pairs_target = [(float(a), float(b)) for a, b in zip(target_shift, losses) if a is not None and b is not None]
        pairs_feature = [(float(a), float(b)) for a, b in zip(feature_shift, losses) if a is not None and b is not None]
        pairs_size = [(float(a), float(b)) for a, b in zip(group_n, losses) if a is not None and b is not None]
        target_rho = spearman([a for a, _ in pairs_target], [b for _, b in pairs_target])
        feature_rho = spearman([a for a, _ in pairs_feature], [b for _, b in pairs_feature])
        size_rho = spearman([a for a, _ in pairs_size], [b for _, b in pairs_size])
        candidates = {
            "feature_shift_associated": abs(feature_rho or 0.0),
            "target_or_class_shift_associated": abs(target_rho or 0.0),
            "group_size_associated": abs(size_rho or 0.0),
        }
        mechanism = max(candidates, key=candidates.get) if candidates else "insufficient_groups"
        summary_rows.append(
            {
                "module": module,
                "descriptor": descriptor,
                "model": model,
                "metric_type": metric_type,
                "n_groups": len(rows),
                "groups_with_feature_shift": sum(1 for row in rows if row["feature_mean_smd"] != ""),
                "median_grouped_loss": fmt(median(usable_loss) if usable_loss else None, 6),
                "median_paired_loss_gap": fmt(median([float(v) for v in gaps if v is not None]) if any(v is not None for v in gaps) else None, 6),
                "spearman_feature_shift_vs_grouped_loss": fmt(feature_rho, 3),
                "spearman_target_shift_vs_grouped_loss": fmt(target_rho, 3),
                "spearman_group_size_vs_grouped_loss": fmt(size_rho, 3),
                "mechanism_fingerprint": mechanism,
            }
        )

    fields = [
        "module",
        "descriptor",
        "model",
        "metric_type",
        "group",
        "group_n",
        "feature_mean_smd",
        "feature_max_smd",
        "n_features_used",
        "target_shift_smd",
        "target_variance_ratio",
        "class_majority_share",
        "class_entropy",
        "random_group_loss",
        "grouped_group_loss",
        "paired_loss_gap",
    ]
    write_csv(FIG_OUT / "r04_groupshift_fingerprint_groups.csv", group_rows, fields)
    summary_fields = [
        "module",
        "descriptor",
        "model",
        "metric_type",
        "n_groups",
        "groups_with_feature_shift",
        "median_grouped_loss",
        "median_paired_loss_gap",
        "spearman_feature_shift_vs_grouped_loss",
        "spearman_target_shift_vs_grouped_loss",
        "spearman_group_size_vs_grouped_loss",
        "mechanism_fingerprint",
    ]
    write_csv(FIG_OUT / "r04_groupshift_fingerprint_summary.csv", summary_rows, summary_fields)

    lines = [
        "# Table S16. R04 group-shift fingerprint diagnostic",
        "",
        "This diagnostic summarizes held-out group stress using target/class shift, group-size stress, and feature standardized-mean-difference scores where public feature matrices can be matched to prediction rows. Spearman associations are exploratory and should not be read causally.",
        "",
        f"- Group-level diagnostic rows: {len(group_rows)}.",
        f"- Descriptor/model summaries: {len(summary_rows)}.",
        "",
        "| Descriptor | model | groups | feature rows | rho(feature, loss) | rho(target/class, loss) | rho(size, loss) | fingerprint |",
        "|---|---|---:|---:|---:|---:|---:|---|",
    ]
    for row in summary_rows:
        if row["model"] == "Random Forest" or row["metric_type"] == "classification":
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['n_groups']} | {row['groups_with_feature_shift']} | {row['spearman_feature_shift_vs_grouped_loss']} | {row['spearman_target_shift_vs_grouped_loss']} | {row['spearman_group_size_vs_grouped_loss']} | {row['mechanism_fingerprint']} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: the atlas contains different stress signatures; section-family extrapolation, source-family classification stress, and ordinary source-reference holdouts should not be collapsed into a single leakage mechanism.",
        ]
    )
    write_text(TABLE_OUT / "r04_groupshift_fingerprint.md", "\n".join(lines) + "\n")


def cvar(values: list[float], tail: float) -> float | None:
    if not values:
        return None
    n_tail = max(1, math.ceil(len(values) * tail))
    return mean(sorted(values, reverse=True)[:n_tail])


def group_tail_risk_and_regret(paired_sets: list[dict[str, object]]) -> None:
    tail_rows: list[dict[str, object]] = []
    regret_rows: list[dict[str, object]] = []
    for data in paired_sets:
        if data["metric_type"] != "regression":
            continue
        labels = [str(x) for x in data["group_labels"]]  # type: ignore[index]
        groups = index_groups(labels)
        if len(groups) < 2:
            continue
        y, pr, pg = numeric_vectors(data)
        if len(y) != len(labels):
            continue
        scale = scale_for(y)
        per_group: list[dict[str, float]] = []
        for group, idx in groups.items():
            yy = [y[i] for i in idx]
            rr = [pr[i] for i in idx]
            gg = [pg[i] for i in idx]
            random_nmae = (mae(yy, rr) or 0.0) / scale
            grouped_nmae = (mae(yy, gg) or 0.0) / scale
            grouped_over = sum(1 for a, b in zip(yy, gg) if b > a) / len(yy)
            random_over = sum(1 for a, b in zip(yy, rr) if b > a) / len(yy)
            grouped_abs = [abs(a - b) / scale for a, b in zip(yy, gg)]
            per_group.append(
                {
                    "n": float(len(idx)),
                    "random_nmae": random_nmae,
                    "grouped_nmae": grouped_nmae,
                    "delta_nmae": grouped_nmae - random_nmae,
                    "random_over": random_over,
                    "grouped_over": grouped_over,
                    "delta_over": grouped_over - random_over,
                    "grouped_p90": percentile(grouped_abs, 0.90) or max(grouped_abs),
                }
            )
        tail_rows.append(
            {
                "module": data["module"],
                "descriptor": data["descriptor"],
                "model": data["model"],
                "n_samples": len(y),
                "n_groups": len(groups),
                "scale_mean_abs_y": fmt(scale, 6),
                "group_balanced_random_nmae": fmt(mean(row["random_nmae"] for row in per_group), 6),
                "group_balanced_grouped_nmae": fmt(mean(row["grouped_nmae"] for row in per_group), 6),
                "delta_group_balanced_nmae": fmt(mean(row["delta_nmae"] for row in per_group), 6),
                "worst_group_grouped_nmae": fmt(max(row["grouped_nmae"] for row in per_group), 6),
                "cvar10_grouped_nmae": fmt(cvar([row["grouped_nmae"] for row in per_group], 0.10), 6),
                "cvar20_grouped_nmae": fmt(cvar([row["grouped_nmae"] for row in per_group], 0.20), 6),
                "cvar20_delta_nmae": fmt(cvar([row["delta_nmae"] for row in per_group], 0.20), 6),
                "group_balanced_random_overprediction": fmt(mean(row["random_over"] for row in per_group), 6),
                "group_balanced_grouped_overprediction": fmt(mean(row["grouped_over"] for row in per_group), 6),
                "cvar20_grouped_p90_abs_error": fmt(cvar([row["grouped_p90"] for row in per_group], 0.20), 6),
            }
        )

        true_order = sorted(range(len(y)), key=lambda i: y[i], reverse=True)
        median_true = median(y)
        for top_share in (0.10, 0.20, 0.30):
            k = max(1, math.ceil(len(y) * top_share))
            true_top = set(true_order[:k])
            true_top_mean = mean(y[i] for i in true_top)
            for split, pred in (("random", pr), ("grouped", pg)):
                pred_top = set(sorted(range(len(y)), key=lambda i: pred[i], reverse=True)[:k])
                overlap = len(true_top & pred_top) / k
                predicted_mean = mean(y[i] for i in pred_top)
                regret = (true_top_mean - predicted_mean) / scale
                false_reassurance = sum(1 for i in pred_top if y[i] < median_true) / k
                rho = spearman(y, pred)
                regret_rows.append(
                    {
                        "module": data["module"],
                        "descriptor": data["descriptor"],
                        "model": data["model"],
                        "split": split,
                        "top_share": fmt(top_share, 2),
                        "top_k": k,
                        "spearman_rank_correlation": fmt(rho, 6),
                        "top_k_overlap_with_true_top": fmt(overlap, 6),
                        "normalized_screening_regret": fmt(regret, 6),
                        "false_reassurance_share": fmt(false_reassurance, 6),
                    }
                )
    tail_fields = [
        "module",
        "descriptor",
        "model",
        "n_samples",
        "n_groups",
        "scale_mean_abs_y",
        "group_balanced_random_nmae",
        "group_balanced_grouped_nmae",
        "delta_group_balanced_nmae",
        "worst_group_grouped_nmae",
        "cvar10_grouped_nmae",
        "cvar20_grouped_nmae",
        "cvar20_delta_nmae",
        "group_balanced_random_overprediction",
        "group_balanced_grouped_overprediction",
        "cvar20_grouped_p90_abs_error",
    ]
    write_csv(FIG_OUT / "r04_group_tail_risk.csv", tail_rows, tail_fields)
    regret_fields = [
        "module",
        "descriptor",
        "model",
        "split",
        "top_share",
        "top_k",
        "spearman_rank_correlation",
        "top_k_overlap_with_true_top",
        "normalized_screening_regret",
        "false_reassurance_share",
    ]
    write_csv(FIG_OUT / "r04_screening_regret.csv", regret_rows, regret_fields)

    lines = [
        "# Table S17. R04 worst-group and CVaR-style decision-risk metrics",
        "",
        "Capacity-like regression rows are summarized over deployment groups. CVaR is computed as the mean of the worst 10% or 20% of held-out groups by grouped normalized MAE; it is a diagnostic tail-risk proxy, not a structural reliability index.",
        "",
        f"- Regression descriptor/model rows: {len(tail_rows)}.",
        "",
        "| Descriptor | model | groups | grouped group-balanced nMAE | worst-group nMAE | CVaR20 nMAE | CVaR20 delta nMAE |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    for row in tail_rows:
        if row["model"] == "Random Forest":
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['n_groups']} | {row['group_balanced_grouped_nmae']} | {row['worst_group_grouped_nmae']} | {row['cvar20_grouped_nmae']} | {row['cvar20_delta_nmae']} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: group-tail metrics expose whether deployment-matched error is concentrated in a small number of sources or structural families, complementing pooled R2 and average nMAE.",
        ]
    )
    write_text(TABLE_OUT / "r04_group_tail_risk.md", "\n".join(lines) + "\n")

    lines = [
        "# Table S18. R04 screening-regret and ranking-stability diagnostics",
        "",
        "For regression rows where larger predicted response is treated as a screening score, this table reports top-k overlap with the true top-k set, normalized regret, false-reassurance share, and Spearman ranking correlation. It is intended for retrofit, experiment-prioritization, or conceptual-design screening interpretations.",
        "",
        f"- Split/model/top-k rows: {len(regret_rows)}.",
        "",
        "| Descriptor | model | split | top share | rank rho | top-k overlap | normalized regret | false reassurance |",
        "|---|---|---|---:|---:|---:|---:|---:|",
    ]
    for row in regret_rows:
        if row["model"] == "Random Forest" and row["top_share"] == "0.20":
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['split']} | {row['top_share']} | {row['spearman_rank_correlation']} | {row['top_k_overlap_with_true_top']} | {row['normalized_screening_regret']} | {row['false_reassurance_share']} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: ranking diagnostics translate validation-target mismatch into screening degradation. They should not be read as project-specific safety decisions.",
        ]
    )
    write_text(TABLE_OUT / "r04_screening_regret.md", "\n".join(lines) + "\n")


def probability_rows(data: dict[str, object], split: str) -> tuple[list[str], list[str], list[dict[str, float]]] | None:
    rows = data["random_rows"] if split == "random" else data["grouped_rows"]  # type: ignore[index]
    if not rows:
        return None
    prob_cols = sorted([key for key in rows[0] if key.startswith("prob_class_")])
    if not prob_cols:
        return None
    y_col = str(data["y_col"])
    p_col = str(data["p_col"])
    y = [normalize_label(row.get(y_col)) for row in rows]  # type: ignore[union-attr]
    pred = [normalize_label(row.get(p_col)) for row in rows]  # type: ignore[union-attr]
    probs: list[dict[str, float]] = []
    for row in rows:  # type: ignore[union-attr]
        pmap: dict[str, float] = {}
        for col in prob_cols:
            label = normalize_label(col.replace("prob_class_", ""))
            value = f(row.get(col))
            if value is not None:
                pmap[label] = value
        probs.append(pmap)
    return y, pred, probs


def expected_calibration_error(y: list[str], pred: list[str], probs: list[dict[str, float]], bins: int = 10) -> tuple[float | None, float | None, float | None]:
    if not y or len(y) != len(pred) or len(y) != len(probs):
        return None, None, None
    confidences = [max(pmap.values()) if pmap else 0.0 for pmap in probs]
    correct = [1.0 if yy == pp else 0.0 for yy, pp in zip(y, pred)]
    ece = 0.0
    for b in range(bins):
        lo = b / bins
        hi = (b + 1) / bins
        idx = [i for i, conf in enumerate(confidences) if (lo <= conf < hi) or (b == bins - 1 and conf <= hi)]
        if not idx:
            continue
        ece += (len(idx) / len(y)) * abs(mean(correct[i] for i in idx) - mean(confidences[i] for i in idx))
    brier_values: list[float] = []
    labels = sorted(set(y) | set().union(*(set(pmap.keys()) for pmap in probs)))
    for yy, pmap in zip(y, probs):
        brier_values.append(sum((pmap.get(label, 0.0) - (1.0 if yy == label else 0.0)) ** 2 for label in labels))
    return ece, mean(confidences), mean(brier_values)


def calibration_diagnostic(paired_sets: list[dict[str, object]]) -> None:
    rows: list[dict[str, object]] = []
    for data in paired_sets:
        if data["metric_type"] == "regression":
            y, pr, pg = numeric_vectors(data)
            if len(y) < 10:
                continue
            scale = scale_for(y)
            random_abs = [abs(a - b) / scale for a, b in zip(y, pr)]
            grouped_abs = [abs(a - b) / scale for a, b in zip(y, pg)]
            for q in (0.80, 0.90, 0.95):
                random_q = percentile(random_abs, q)
                grouped_q = percentile(grouped_abs, q)
                if random_q is None or grouped_q is None:
                    continue
                rows.append(
                    {
                        "module": data["module"],
                        "descriptor": data["descriptor"],
                        "model": data["model"],
                        "metric_type": data["metric_type"],
                        "calibration_basis": f"absolute residual q{int(q*100)}",
                        "random_calibrated_width_norm": fmt(random_q, 6),
                        "grouped_calibrated_width_norm": fmt(grouped_q, 6),
                        "width_inflation_grouped_over_random": fmt(grouped_q / random_q if random_q > EPS else None, 6),
                        "random_interval_coverage_on_random": fmt(sum(1 for value in random_abs if value <= random_q) / len(random_abs), 6),
                        "random_interval_coverage_on_grouped": fmt(sum(1 for value in grouped_abs if value <= random_q) / len(grouped_abs), 6),
                        "coverage_shortfall_under_grouped": fmt(q - (sum(1 for value in grouped_abs if value <= random_q) / len(grouped_abs)), 6),
                        "ece": "",
                        "mean_confidence": "",
                        "brier_score": "",
                    }
                )
        else:
            for split in ("random", "grouped"):
                values = probability_rows(data, split)
                if values is None:
                    continue
                y_cls, pred_cls, probs = values
                ece, confidence, brier = expected_calibration_error(y_cls, pred_cls, probs)
                rows.append(
                    {
                        "module": data["module"],
                        "descriptor": data["descriptor"],
                        "model": data["model"],
                        "metric_type": data["metric_type"],
                        "calibration_basis": f"{split} probability calibration",
                        "random_calibrated_width_norm": "",
                        "grouped_calibrated_width_norm": "",
                        "width_inflation_grouped_over_random": "",
                        "random_interval_coverage_on_random": "",
                        "random_interval_coverage_on_grouped": "",
                        "coverage_shortfall_under_grouped": "",
                        "ece": fmt(ece, 6),
                        "mean_confidence": fmt(confidence, 6),
                        "brier_score": fmt(brier, 6),
                    }
                )
    fields = [
        "module",
        "descriptor",
        "model",
        "metric_type",
        "calibration_basis",
        "random_calibrated_width_norm",
        "grouped_calibrated_width_norm",
        "width_inflation_grouped_over_random",
        "random_interval_coverage_on_random",
        "random_interval_coverage_on_grouped",
        "coverage_shortfall_under_grouped",
        "ece",
        "mean_confidence",
        "brier_score",
    ]
    write_csv(FIG_OUT / "r04_calibration_diagnostic.csv", rows, fields)

    lines = [
        "# Table S19. R04 source-aware calibration diagnostic",
        "",
        "Regression rows use random-validation residual quantiles as simple absolute-error intervals and test their empirical coverage under grouped predictions. Classification rows with exported probabilities report ECE, mean confidence, and Brier score. These diagnostics do not replace formal uncertainty calibration or structural reliability analysis.",
        "",
        f"- Calibration rows: {len(rows)}.",
        "",
        "| Descriptor | model | basis | width inflation | grouped coverage using random width | shortfall/ECE |",
        "|---|---|---|---:|---:|---:|",
    ]
    for row in rows:
        if row["model"] == "Random Forest" and ("q90" in str(row["calibration_basis"]) or "probability" in str(row["calibration_basis"])):
            shortfall = row["coverage_shortfall_under_grouped"] if row["coverage_shortfall_under_grouped"] != "" else row["ece"]
            coverage = row["random_interval_coverage_on_grouped"] if row["random_interval_coverage_on_grouped"] != "" else row["mean_confidence"]
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['calibration_basis']} | {row['width_inflation_grouped_over_random']} | {coverage} | {shortfall} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: intervals or probability confidence calibrated under random validation can be over-optimistic under deployment-matched grouped holdout, reinforcing the need to export uncertainty diagnostics beside headline scores.",
        ]
    )
    write_text(TABLE_OUT / "r04_calibration_diagnostic.md", "\n".join(lines) + "\n")


def source_balanced_scoring_probe(paired_sets: list[dict[str, object]]) -> None:
    rows: list[dict[str, object]] = []
    for data in paired_sets:
        labels = [str(x) for x in data["group_labels"]]  # type: ignore[index]
        groups = index_groups(labels)
        if len(groups) < 2:
            continue
        if data["metric_type"] == "regression":
            y, pr, pg = numeric_vectors(data)
            if len(y) != len(labels):
                continue
            scale = scale_for(y)
            random_loss = [abs(a - b) / scale for a, b in zip(y, pr)]
            grouped_loss = [abs(a - b) / scale for a, b in zip(y, pg)]
            pooled_random = mean(random_loss)
            pooled_grouped = mean(grouped_loss)
            source_random = group_balanced_mean(random_loss, labels)
            source_grouped = group_balanced_mean(grouped_loss, labels)
            basis = "normalized absolute error"
        else:
            y_cls, pr_cls, pg_cls = class_vectors(data)
            if len(y_cls) != len(labels):
                continue
            random_loss = [0.0 if yy == pp else 1.0 for yy, pp in zip(y_cls, pr_cls)]
            grouped_loss = [0.0 if yy == pp else 1.0 for yy, pp in zip(y_cls, pg_cls)]
            pooled_random = mean(random_loss)
            pooled_grouped = mean(grouped_loss)
            source_random = group_balanced_mean(random_loss, labels)
            source_grouped = group_balanced_mean(grouped_loss, labels)
            basis = "classification error"
        if source_random is None or source_grouped is None:
            continue
        pooled_delta = pooled_grouped - pooled_random
        source_delta = source_grouped - source_random
        rows.append(
            {
                "module": data["module"],
                "descriptor": data["descriptor"],
                "model": data["model"],
                "metric_type": data["metric_type"],
                "basis": basis,
                "n_samples": len(labels),
                "n_groups": len(groups),
                "pooled_random_loss": fmt(pooled_random, 6),
                "pooled_grouped_loss": fmt(pooled_grouped, 6),
                "pooled_loss_delta_grouped_minus_random": fmt(pooled_delta, 6),
                "source_balanced_random_loss": fmt(source_random, 6),
                "source_balanced_grouped_loss": fmt(source_grouped, 6),
                "source_balanced_loss_delta_grouped_minus_random": fmt(source_delta, 6),
                "source_balancing_changes_delta_by": fmt(source_delta - pooled_delta, 6),
                "probe_interpretation": "source_weighting_exposes_more_degradation"
                if source_delta - pooled_delta > 0.05
                else "source_weighting_similar_to_pooled"
                if abs(source_delta - pooled_delta) <= 0.05
                else "source_weighting_exposes_less_degradation",
            }
        )
    fields = [
        "module",
        "descriptor",
        "model",
        "metric_type",
        "basis",
        "n_samples",
        "n_groups",
        "pooled_random_loss",
        "pooled_grouped_loss",
        "pooled_loss_delta_grouped_minus_random",
        "source_balanced_random_loss",
        "source_balanced_grouped_loss",
        "source_balanced_loss_delta_grouped_minus_random",
        "source_balancing_changes_delta_by",
        "probe_interpretation",
    ]
    write_csv(FIG_OUT / "r04_source_balanced_scoring_probe.csv", rows, fields)

    lines = [
        "# Table S20. R04 source-balanced scoring probe",
        "",
        "This table applies equal source/group weights at the scoring stage. It is a conservative mitigation-design probe: it indicates where source-balanced ERM would be worth testing, but it does not claim to be a retrained source-balanced ERM benchmark.",
        "",
        f"- Descriptor/model rows: {len(rows)}.",
        "",
        "| Descriptor | model | basis | pooled delta | source-balanced delta | change | interpretation |",
        "|---|---|---|---:|---:|---:|---|",
    ]
    for row in rows:
        if row["model"] == "Random Forest" or row["metric_type"] == "classification":
            lines.append(
                f"| {row['descriptor']} | {row['model']} | {row['basis']} | {row['pooled_loss_delta_grouped_minus_random']} | {row['source_balanced_loss_delta_grouped_minus_random']} | {row['source_balancing_changes_delta_by']} | {row['probe_interpretation']} |"
            )
    lines.extend(
        [
            "",
            "Interpretation: equal-source scoring identifies rows where deployment groups, not only samples, should drive mitigation and reporting. A retrained source-balanced ERM baseline remains a separate extension requiring fold-level retraining.",
        ]
    )
    write_text(TABLE_OUT / "r04_source_balanced_scoring_probe.md", "\n".join(lines) + "\n")


def main() -> None:
    paired_sets = collect_prediction_pairs()
    topology_pseudogroup_null(paired_sets)
    group_shift_fingerprint(paired_sets)
    group_tail_risk_and_regret(paired_sets)
    calibration_diagnostic(paired_sets)
    source_balanced_scoring_probe(paired_sets)
    outputs = [
        FIG_OUT / "r04_topology_pseudogroup_null.csv",
        FIG_OUT / "r04_groupshift_fingerprint_groups.csv",
        FIG_OUT / "r04_groupshift_fingerprint_summary.csv",
        FIG_OUT / "r04_group_tail_risk.csv",
        FIG_OUT / "r04_screening_regret.csv",
        FIG_OUT / "r04_calibration_diagnostic.csv",
        FIG_OUT / "r04_source_balanced_scoring_probe.csv",
        TABLE_OUT / "r04_topology_pseudogroup_null.md",
        TABLE_OUT / "r04_groupshift_fingerprint.md",
        TABLE_OUT / "r04_group_tail_risk.md",
        TABLE_OUT / "r04_screening_regret.md",
        TABLE_OUT / "r04_calibration_diagnostic.md",
        TABLE_OUT / "r04_source_balanced_scoring_probe.md",
    ]
    for path in outputs:
        print(f"[OK] {path}")


if __name__ == "__main__":
    main()
