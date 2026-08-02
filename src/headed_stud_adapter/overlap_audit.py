"""Outcome-blind source overlap audit for the headed-stud parent family."""

from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"
PARENTS = {
    "nwc242": "development_parent",
    "deck464": "primary_external_parent",
    "lwc90": "near_domain_control",
}


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def source_codes(value: object) -> list[str]:
    return re.findall(r"\[\s*\d+\s*\]", str(value or ""))


def author_year_key(author: object, citation: object) -> str | None:
    author_token = normalize(author).split(" ")[0] if normalize(author) else ""
    years = re.findall(r"(?:19|20)\d{2}", str(citation or ""))
    return f"{author_token}:{years[-1]}" if author_token and years else None


def archive_workbook(short_name: str) -> openpyxl.Workbook:
    archive = RAW / f"{short_name}.zip"
    with zipfile.ZipFile(archive) as outer:
        names = [name for name in outer.namelist() if name.lower().endswith(".xlsx")]
        if len(names) != 1:
            raise RuntimeError(f"expected one XLSX in {archive.name}, got {names}")
        data = outer.read(names[0])
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)


def parse_parent(short_name: str) -> list[dict]:
    workbook = archive_workbook(short_name)
    notes = workbook["Notes"]
    citation_map: dict[str, str] = {}
    for row in notes.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0] is not None and row[1] is not None:
            for code in source_codes(row[0]):
                citation_map[code] = str(row[1]).strip()

    sheet = workbook["Database"]
    headers = {
        int(cell.column): str(cell.value).strip()
        for cell in next(sheet.iter_rows(min_row=6, max_row=6))
        if cell.value is not None
    }
    target_columns = [
        column
        for column, header in headers.items()
        if normalize(header) in {"pem kn", "measured resistance pem kn"}
    ]
    if target_columns != [5]:
        raise RuntimeError(
            f"unexpected target column for {short_name}: {target_columns}"
        )

    rows = []
    for excel_row in range(7, sheet.max_row + 1):
        specimen = sheet.cell(excel_row, 2).value
        author = sheet.cell(excel_row, 3).value
        raw_code = sheet.cell(excel_row, 4).value
        if specimen is None and author is None and raw_code is None:
            continue
        codes = source_codes(raw_code)
        citations = [citation_map.get(code, "") for code in codes]
        citation_keys = [normalize(citation) for citation in citations if citation]
        exact_keys = [
            hashlib.sha256(key.encode("utf-8")).hexdigest()[:20]
            for key in citation_keys
        ]
        probable_keys = [
            key
            for key in (
                author_year_key(author, citation) for citation in citations
            )
            if key
        ]
        rows.append(
            {
                "parent": short_name,
                "role": PARENTS[short_name],
                "excel_row": excel_row,
                "specimen_reference": str(specimen or "").strip(),
                "source_author": str(author or "").strip(),
                "source_codes": "|".join(codes),
                "source_citations": " || ".join(citations),
                "exact_source_keys": "|".join(exact_keys),
                "author_year_keys": "|".join(probable_keys),
                "source_recovered": bool(exact_keys or probable_keys),
                "target_value_accessed": False,
            }
        )
    workbook.close()
    return rows


def key_set(rows: list[dict], field: str) -> set[str]:
    return {
        token
        for row in rows
        for token in str(row[field]).split("|")
        if token
    }


def main() -> None:
    all_rows = {name: parse_parent(name) for name in PARENTS}
    development_exact = key_set(all_rows["nwc242"], "exact_source_keys")
    development_probable = key_set(all_rows["nwc242"], "author_year_keys")
    flat_rows = []
    for parent, rows in all_rows.items():
        for row in rows:
            exact = bool(
                key_set([row], "exact_source_keys") & development_exact
            ) if parent != "nwc242" else False
            probable = bool(
                key_set([row], "author_year_keys") & development_probable
            ) if parent != "nwc242" else False
            row["exact_development_source_overlap"] = exact
            row["probable_author_year_overlap"] = probable and not exact
            row["quarantine_candidate"] = exact or probable
            flat_rows.append(row)

    frame = pd.DataFrame(flat_rows)
    frame.to_csv(HERE / "source_audit.csv", index=False, encoding="utf-8")
    summary = []
    for parent, group in frame.groupby("parent", sort=False):
        source_tokens = [
            token
            for value in group["exact_source_keys"]
            for token in str(value).split("|")
            if token
        ]
        retained = group.loc[~group["quarantine_candidate"]]
        retained_tokens = [
            token
            for value in retained["exact_source_keys"]
            for token in str(value).split("|")
            if token
        ]
        counts = Counter(retained_tokens)
        summary.append(
            {
                "parent": parent,
                "role": group["role"].iloc[0],
                "n_rows": int(len(group)),
                "source_recovery_rate": float(group["source_recovered"].mean()),
                "n_exact_source_keys": len(set(source_tokens)),
                "n_exact_overlap_rows": int(
                    group["exact_development_source_overlap"].sum()
                ),
                "n_probable_overlap_rows": int(
                    group["probable_author_year_overlap"].sum()
                ),
                "n_quarantine_candidates": int(
                    group["quarantine_candidate"].sum()
                ),
                "n_retained_rows_pending_manual_review": int(len(retained)),
                "n_retained_exact_source_keys": len(set(retained_tokens)),
                "n_retained_singleton_source_keys": int(
                    sum(count == 1 for count in counts.values())
                ),
            }
        )
    payload = {
        "schema_version": "r30-headed-stud-source-overlap-v1",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "target_values_accessed": False,
        "quarantine_status": "candidates_pending_blind_manual_citation_review",
        "summary": summary,
        "source_audit_sha256": hashlib.sha256(
            (HERE / "source_audit.csv").read_bytes()
        ).hexdigest(),
    }
    (HERE / "overlap.json").write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(payload, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
