"""Unseal only the frozen target and build private canonical study tables."""

from __future__ import annotations

import hashlib
import io
import json
import math
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"
DERIVED = HERE / "derived_private"
ADAPTER = HERE / "adapter_freeze_v1.json"
PLAN = HERE / "evaluation_plan_v1.md"
AUDIT = HERE / "source_audit_adjudicated.csv"
EXPECTED_ADAPTER_SHA256 = "42ac1a0bed31307884ce19bc38eb629cff198cd67c06b6ea1347af1f3d919515"
EXPECTED_PLAN_SHA256 = "9dabb8178a8f1d96fe6278adb8ce4e38d9e1e0e8dc7699c12a4a001b00d09a76"
FEATURES = [
    "f_cm_MPa",
    "E_cm_MPa",
    "f_um_MPa",
    "d_m_mm",
    "d_dom_mm",
    "h_wm_mm",
    "h_scm_mm",
]
EXCEL_CONFIG = {
    "nwc242": {
        "archive": "nwc242.zip",
        "columns": [7, 13, 14, 15, 16, 17, 19],
        "topology": "solid",
        "concrete_family": "normal_weight",
        "role": "development_parent",
    },
    "deck464": {
        "archive": "deck464.zip",
        "columns": [26, 32, 33, 35, 36, 37, 39],
        "topology": "profiled_deck",
        "concrete_family": "normal_weight",
        "role": "primary_external_parent",
    },
    "lwc90": {
        "archive": "lwc90.zip",
        "columns": [7, 13, 15, 16, 17, 18, 20],
        "topology": "solid",
        "concrete_family": "lightweight",
        "role": "near_domain_control",
    },
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_freeze() -> dict:
    if sha256(ADAPTER) != EXPECTED_ADAPTER_SHA256:
        raise RuntimeError("adapter freeze changed after target-unseal authorization")
    if sha256(PLAN) != EXPECTED_PLAN_SHA256:
        raise RuntimeError("evaluation plan changed after target-unseal authorization")
    freeze = json.loads(ADAPTER.read_text(encoding="utf-8"))
    for filename, expected in {
        "freeze_v1.json": freeze["frozen_input_hashes"]["parent_registry"],
        "download_manifest.json": freeze["frozen_input_hashes"]["download_manifest"],
        "schema_scan_v1.json": freeze["frozen_input_hashes"]["schema_scan"],
        "source_audit_adjudicated.csv": freeze["frozen_input_hashes"]["adjudicated_source_rows"],
        "overlap_adjudicated.json": freeze["frozen_input_hashes"]["adjudication_summary"],
    }.items():
        if sha256(HERE / filename) != expected:
            raise RuntimeError(f"frozen input hash mismatch: {filename}")
    return freeze


def source_lookup() -> dict[tuple[str, int], dict]:
    frame = pd.read_csv(AUDIT).fillna("")
    frame["final_quarantine"] = (
        frame["final_quarantine"].astype(str).str.lower() == "true"
    )
    return {
        (str(row.parent), int(row.excel_row)): row._asdict()
        for row in frame.itertuples(index=False)
    }


def workbook_from_archive(archive_name: str) -> openpyxl.Workbook:
    with zipfile.ZipFile(RAW / archive_name) as outer:
        members = [name for name in outer.namelist() if name.lower().endswith(".xlsx")]
        if len(members) != 1:
            raise RuntimeError(f"expected one XLSX in {archive_name}")
        data = outer.read(members[0])
    return openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)


def build_excel_parent(parent: str, lookup: dict[tuple[str, int], dict]) -> tuple[pd.DataFrame, dict]:
    config = EXCEL_CONFIG[parent]
    workbook = workbook_from_archive(config["archive"])
    sheet = workbook["Database"]
    rows = []
    for excel_row in range(7, sheet.max_row + 1):
        audit = lookup.get((parent, excel_row))
        if audit is None:
            continue
        exact_keys = str(audit["exact_source_keys"]).strip()
        source_group = exact_keys.replace("|", "+") or "__UNRESOLVED_SOURCE__"
        values = [sheet.cell(excel_row, column).value for column in config["columns"]]
        row = {
            "parent": parent,
            "role": config["role"],
            "row_id": f"{parent}:xlsx:{excel_row}",
            "original_row": excel_row,
            "specimen_reference": str(audit["specimen_reference"]),
            "source_group": source_group,
            "source_recovered": bool(audit["source_recovered"]),
            "source_quarantine": bool(audit["final_quarantine"]),
            "slab_topology": config["topology"],
            "concrete_family": config["concrete_family"],
            "P_em_kN": sheet.cell(excel_row, 5).value,
        }
        row.update(dict(zip(FEATURES, values)))
        rows.append(row)
    workbook.close()
    raw = pd.DataFrame(rows)
    if parent == "nwc242" and len(raw) != 242:
        raise RuntimeError(f"NWC row-count mismatch: {len(raw)}")
    if parent == "deck464" and len(raw) != 464:
        raise RuntimeError(f"deck row-count mismatch: {len(raw)}")
    if parent == "lwc90" and len(raw) != 90:
        raise RuntimeError(f"LWC row-count mismatch: {len(raw)}")
    target = pd.to_numeric(raw["P_em_kN"], errors="coerce")
    raw_target_range = [float(target.min()), float(target.max())]
    if parent == "nwc242" and not all(
        math.isclose(actual, expected, rel_tol=0, abs_tol=0.02)
        for actual, expected in zip(raw_target_range, [61.83, 318.90])
    ):
        raise RuntimeError(f"NWC published target range mismatch: {raw_target_range}")
    if parent == "deck464" and not all(
        math.isclose(actual, expected, rel_tol=0, abs_tol=0.02)
        for actual, expected in zip(raw_target_range, [35.00, 130.00])
    ):
        raise RuntimeError(f"deck published target range mismatch: {raw_target_range}")

    source_removed = int(raw["source_quarantine"].sum())
    retained = raw.loc[~raw["source_quarantine"]].copy()
    numeric = FEATURES + ["P_em_kN"]
    for column in numeric:
        retained[column] = pd.to_numeric(retained[column], errors="coerce")
    complete = retained[numeric].notna().all(axis=1)
    missing_removed = int((~complete).sum())
    retained = retained.loc[complete].copy()
    positive = retained["P_em_kN"] > 0
    nonpositive_removed = int((~positive).sum())
    retained = retained.loc[positive].reset_index(drop=True)
    summary = {
        "parent": parent,
        "n_raw": int(len(raw)),
        "n_source_quarantine": source_removed,
        "n_missing_target_or_feature": missing_removed,
        "n_nonpositive_target": nonpositive_removed,
        "n_retained": int(len(retained)),
        "raw_target_range_kN": raw_target_range,
        "retained_target_range_kN": [
            float(retained["P_em_kN"].min()),
            float(retained["P_em_kN"].max()),
        ],
        "n_retained_exact_source_groups": int(
            retained.loc[
                retained["source_group"] != "__UNRESOLVED_SOURCE__",
                "source_group",
            ].nunique()
        ),
        "n_unresolved_source_rows": int(
            (retained["source_group"] == "__UNRESOLVED_SOURCE__").sum()
        ),
    }
    return retained, summary


def build_rac() -> tuple[pd.DataFrame, dict]:
    with zipfile.ZipFile(RAW / "rac27.zip") as outer:
        data = outer.read("Machine learning/solid_RAC_27points.csv")
    raw = pd.read_csv(io.BytesIO(data))
    mapping = {
        "P_em": "P_em_kN",
        "f_cm": "f_cm_MPa",
        "E_cm": "E_cm_MPa",
        "f_um": "f_um_MPa",
        "d_m": "d_m_mm",
        "d_dom": "d_dom_mm",
        "h_wm": "h_wm_mm",
        "h_m": "h_scm_mm",
    }
    if len(raw) != 27 or set(raw["RA_%"].unique()) != {20, 40, 60}:
        raise RuntimeError("RAC frozen subset does not contain the expected 27 positive-RCA rows")
    frame = raw.rename(columns=mapping)
    frame.insert(0, "parent", "rac27")
    frame.insert(1, "role", "small_independent_experiment_confirmation")
    frame.insert(2, "row_id", [f"rac27:csv:{index + 2}" for index in range(len(frame))])
    frame.insert(3, "original_row", range(2, len(frame) + 2))
    frame.insert(
        4,
        "specimen_reference",
        [
            f"RAC{int(row['RA_%']):02d}_fc{row['f_cm_MPa']:.1f}_d{row['d_m_mm']:.1f}"
            for _, row in frame.iterrows()
        ],
    )
    frame.insert(5, "source_group", "basrah_2022_single_programme")
    frame.insert(6, "source_recovered", True)
    frame.insert(7, "source_quarantine", False)
    frame.insert(8, "slab_topology", "solid")
    frame.insert(9, "concrete_family", "recycled_aggregate_scc")
    numeric = FEATURES + ["P_em_kN"]
    for column in numeric:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    complete = frame[numeric].notna().all(axis=1)
    positive = frame["P_em_kN"] > 0
    retained = frame.loc[complete & positive].reset_index(drop=True)
    if len(retained) != 27:
        raise RuntimeError("RAC rows failed the frozen complete/positive gate")
    summary = {
        "parent": "rac27",
        "n_raw": 27,
        "n_source_quarantine": 0,
        "n_missing_target_or_feature": int((~complete).sum()),
        "n_nonpositive_target": int((~positive).sum()),
        "n_retained": int(len(retained)),
        "raw_target_range_kN": [
            float(frame["P_em_kN"].min()),
            float(frame["P_em_kN"].max()),
        ],
        "retained_target_range_kN": [
            float(retained["P_em_kN"].min()),
            float(retained["P_em_kN"].max()),
        ],
        "n_retained_exact_source_groups": 1,
        "n_unresolved_source_rows": 0,
    }
    return retained, summary


def main() -> None:
    freeze = verify_freeze()
    lookup = source_lookup()
    DERIVED.mkdir(parents=True, exist_ok=True)
    tables: dict[str, pd.DataFrame] = {}
    summaries = []
    for parent in EXCEL_CONFIG:
        table, summary = build_excel_parent(parent, lookup)
        tables[parent] = table
        summaries.append(summary)
    tables["rac27"], rac_summary = build_rac()
    summaries.append(rac_summary)
    outputs = []
    for parent, table in tables.items():
        path = DERIVED / f"{parent}.csv"
        table.to_csv(path, index=False, encoding="utf-8")
        outputs.append(
            {
                "parent": parent,
                "path": str(path.relative_to(HERE)),
                "rows": int(len(table)),
                "sha256": sha256(path),
                "private_third_party_derived_data": True,
            }
        )
    manifest = {
        "schema_version": "r30-headed-stud-adapter-build-v1",
        "target_unsealed_utc": datetime.now(timezone.utc).isoformat(),
        "estimand": freeze["estimand"],
        "adapter_freeze_sha256": EXPECTED_ADAPTER_SHA256,
        "evaluation_plan_sha256": EXPECTED_PLAN_SHA256,
        "summaries": summaries,
        "outputs": outputs,
        "target_access_deviation": None,
        "public_release": "adapters and aggregate summaries only; canonical row tables stay private",
    }
    (HERE / "adapter_manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
